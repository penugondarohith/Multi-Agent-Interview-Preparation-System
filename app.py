import os
import sys
import tempfile
from pathlib import Path
from types import ModuleType

# Mock chromadb to avoid Pydantic V1 type-inference compatibility issues in Python 3.14+
class DummyBase:
    def __class_getitem__(cls, item):
        return cls

    @classmethod
    def __get_pydantic_core_schema__(cls, source_type, handler):
        from pydantic_core import core_schema
        return core_schema.any_schema()

class MockModule(ModuleType):
    def __init__(self, name):
        super().__init__(name)
        self.__path__ = []

    def __getattr__(self, name):
        if name.startswith('__'):
            raise AttributeError(name)
        class DynamicMockClass(DummyBase):
            pass
        DynamicMockClass.__name__ = name
        return DynamicMockClass

class ChromadbMockFinder:
    def find_spec(self, fullname, path, target=None):
        if fullname.startswith('chromadb'):
            from importlib.machinery import ModuleSpec
            return ModuleSpec(fullname, self)
        return None

    def create_module(self, spec):
        return MockModule(spec.name)

    def exec_module(self, module):
        pass

sys.meta_path.insert(0, ChromadbMockFinder())

# Disable telemetry and OpenTelemetry tracer overrides before other imports to prevent crashes on subsequent runs
os.environ["CREWAI_DISABLE_TELEMETRY"] = "true"
os.environ["OTEL_SDK_DISABLED"] = "true"

import nest_asyncio
nest_asyncio.apply()

import crewai.llms.cache as _crewai_cache
# Monkey-patch to prevent injection of unsupported cache_breakpoint parameter
_crewai_cache.mark_cache_breakpoint = lambda msg: msg

import asyncio
import pdfplumber
import streamlit as st
import time
from docx import Document
from crewai import Agent, Task, Crew, LLM
from concurrent.futures import ThreadPoolExecutor, as_completed

# Force standard asyncio event loop policy to avoid issues with uvloop on Streamlit Cloud (Linux)
try:
    asyncio.set_event_loop_policy(asyncio.DefaultEventLoopPolicy())
except Exception:
    pass

try:
    loop = asyncio.new_event_loop()
    asyncio.set_event_loop(loop)
except Exception:
    loop = asyncio.get_event_loop()


def get_temp_storage_dir():
    storage_dir = Path(tempfile.gettempdir()) / "resume_interview_prep"
    storage_dir.mkdir(exist_ok=True)
    return storage_dir


def save_uploaded_file(uploaded_file):
    suffix = os.path.splitext(uploaded_file.name)[1]
    storage_dir = get_temp_storage_dir()
    with tempfile.NamedTemporaryFile(dir=storage_dir, delete=False, suffix=suffix) as tmp:
        tmp.write(uploaded_file.getvalue())
        return tmp.name

st.set_page_config(page_title="Resume Interview Prep", layout="wide")

st.markdown(
    """
    <style>
    .agent-card {
      min-width: 260px;
      max-width: 350px;

      width: 100%;
      height: auto;

      min-height: 220px;

      padding: 70px;
      margin: 30px 12px 24px 12px !important; /* Visible space between neighboring cards */
      box-sizing: border-box;

      border-radius: 20px;

      display: flex;
      flex-direction: column;
      align-items: center;

      word-wrap: break-word;
      overflow-wrap: break-word;
    }

    /* Individual card column behavior inside each row */
    div[data-testid="column"]:has(.agent-card) {
      flex: 1 1 320px;
      min-width: 320px !important;
      max-width: 380px;
      margin: 0 0 8px 0;
      padding: 0 12px !important; /* Extra side padding so the gap appears between cards */
      box-sizing: border-box;
    }

    /* Keep cards wrapping to the next row when space is limited */
    div[data-testid="stHorizontalBlock"]:has(.agent-card) {
      display: flex !important;
      flex-wrap: wrap !important;
      gap: 0 24px !important; /* Strong horizontal gap between individual cards in the same row */
      justify-content: space-between;
      align-items: stretch;
    }
    </style>
    """,
    unsafe_allow_html=True,
)

st.title("Resume-to-Interview Preparation Assistant")
st.caption("Upload your resume, enter the role/company, and generate an interactive interview prep report.")
st.info("Use an OpenRouter API key (the key usually starts with 'sk-or-...').")
st.markdown("Create or view your OpenRouter key: https://openrouter.ai/keys")

OPENROUTER_API_KEY = st.text_input(
    "OpenRouter API Key",
    type="password",
    value=os.getenv("OPENROUTER_API_KEY", ""),
    help="Paste your OpenRouter API key here.",
)
if OPENROUTER_API_KEY:
    os.environ["OPENROUTER_API_KEY"] = OPENROUTER_API_KEY

# Model Selection
model_option = st.selectbox(
    "Select OpenRouter Model",
    options=[
        "openrouter/meta-llama/llama-3.3-70b-instruct",
        "openrouter/google/gemini-2.5-flash",
        "openrouter/deepseek/deepseek-chat",
        "openrouter/anthropic/claude-3.5-sonnet",
    ],
    index=0,
    help="Select the OpenRouter model you want to use for the analysis.",
)
uploaded_file = st.file_uploader(
    "Upload Resume (PDF or DOCX)",
    type=["pdf", "docx"],
    accept_multiple_files=False,
)
job_role = st.text_input("Target Role", placeholder="e.g. Python Developer")
company_name = st.text_input("Company Name", placeholder="e.g. Acme Corp")
job_description = st.text_area(
    "Job Description",
    placeholder="Paste the JD here...",
    height=180,
)

if not uploaded_file:
    st.info("Upload a PDF or DOCX resume to begin.")
    st.stop()
if not job_role or not company_name or not job_description:
    st.warning("Please fill in the role, company name, and job description before generating the report.")
    st.stop()
if not OPENROUTER_API_KEY:
    st.error("An OpenRouter API key is required to generate the report.")
    st.stop()

filename = save_uploaded_file(uploaded_file)
st.success(f"Resume selected: {uploaded_file.name}")

def extract_pdf_text(file_path):
    text = ""
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n"
    return text

def extract_docx_text(file_path):
    doc = Document(file_path)
    return "\n".join(
        para.text for para in doc.paragraphs
    )

if filename.endswith(".pdf"):
    resume_text = extract_pdf_text(filename)
elif filename.endswith(".docx"):
    resume_text = extract_docx_text(filename)
else:
    raise Exception(
        "Only PDF and DOCX files are supported."
    )

# Change Detection: reset generation state if any input parameters change
current_inputs = {
    "uploaded_filename": uploaded_file.name if uploaded_file else "",
    "job_role": job_role,
    "company_name": company_name,
    "job_description": job_description,
    "model_option": model_option,
}

if "previous_inputs" not in st.session_state:
    st.session_state.previous_inputs = current_inputs

if st.session_state.previous_inputs != current_inputs:
    st.session_state.previous_inputs = current_inputs
    st.session_state.incremental_generating = False
    # Clear specific outputs from session state
    for key in [
        "resume_output", "company_output", "jd_match_output", 
        "resume_enhancement_output", "hr_output", "technical_output", 
        "final_output", "mock_interview_output", "resume_creation_output"
    ]:
        if key in st.session_state:
            del st.session_state[key]

def is_quota_exhausted_error(exc):
    exc_str = str(exc).lower()
    return "quota" in exc_str or "exhausted" in exc_str or "429" in exc_str or "rate limit" in exc_str or "credit" in exc_str or "insufficient" in exc_str

def get_llm():
    return LLM(
        model=model_option,
        api_key=OPENROUTER_API_KEY,
        temperature=0.2,
    )

# Individual Step Execution Functions
def run_step1(llm):
    agent = Agent(
        role="Resume Analyzer",
        goal="Analyze resume and ATS readiness.",
        backstory="Senior ATS Consultant.",
        llm=llm,
        verbose=False,
        allow_delegation=False
    )
    task = Task(
        description=f"""
    Analyze the resume and provide a short, simple, and concise ATS analysis:
    Resume:
    {resume_text}
    Role:
    {job_role}
    JD:
    {job_description}
    Provide (keep descriptions short):
    1. ATS Score & Match Score (out of 100)
    2. Top 3 Strengths (brief)
    3. Top 3 Weaknesses (brief)
    4. Missing Skills & Gaps (brief list)
    5. Recommended Certifications & Projects (brief)
    """,
        expected_output="ATS Analysis Report",
        agent=agent
    )
    crew = Crew(agents=[agent], tasks=[task], verbose=False)
    return str(crew.kickoff())

def run_step2(llm):
    agent = Agent(
        role="Company Research Specialist",
        goal="Research company hiring expectations.",
        backstory="Recruitment Research Expert.",
        llm=llm,
        verbose=False,
        allow_delegation=False
    )
    task = Task(
        description=f"""
    Research the company and provide a short, simple, and concise overview:
    Company:
    {company_name}
    Role:
    {job_role}
    Provide (keep each section to a few brief bullet points):
    1. Company Overview (max 3 sentences)
    2. Technology Stack (key tech only)
    3. Hiring Process (brief summary)
    4. Required Skills (top 5 skills)
    5. Interview Pattern (brief)
    6. Preparation Tips (top 3 tips)
    """,
        expected_output="Company Research Report",
        agent=agent
    )
    crew = Crew(agents=[agent], tasks=[task], verbose=False)
    return str(crew.kickoff())

def run_step3(llm):
    agent = Agent(
        role="JD Match Analyzer",
        goal="Compare resume against the job description.",
        backstory="ATS optimization expert.",
        llm=llm,
        verbose=False,
        allow_delegation=False
    )
    task = Task(
        description=f"""
        Analyze JD match details and keep it short, simple, and concise:
        Resume:
        {resume_text}
        Job Description:
        {job_description}
        Analyze (in brief bullet points):
        1. ATS Match %
        2. Top Missing Keywords
        3. Top Missing Skills
        4. Recommended Improvements
        """,
        expected_output="JD Match Analysis",
        agent=agent
    )
    crew = Crew(agents=[agent], tasks=[task], verbose=False)
    return str(crew.kickoff())

def run_step4(llm):
    agent = Agent(
        role="Resume Enhancement Expert",
        goal="Improve ATS score and role alignment.",
        backstory="Professional Resume Writer.",
        llm=llm,
        verbose=False,
        allow_delegation=False
    )
    task = Task(
        description=f"""
    Provide short, simple, and concise resume enhancement suggestions:
    Role:
    {job_role}
    JD:
    {job_description}
    Original Resume:
    {resume_text}

    Provide (keep suggestions short and actionable):
    1. Improved Summary (brief)
    2. Top skills to add
    3. Top project enhancement suggestions
    4. Top missing keywords
    5. Top rewrite suggestions
    """,
        expected_output="Resume Enhancement Suggestions",
        agent=agent
    )
    crew = Crew(agents=[agent], tasks=[task], verbose=False)
    return str(crew.kickoff())

def run_step5(llm):
    agent = Agent(
        role="HR Interviewer",
        goal="Generate company-specific HR questions.",
        backstory="Senior HR Manager.",
        llm=llm,
        verbose=False,
        allow_delegation=False
    )
    task = Task(
        description=f"""
    Generate a short, simple, and concise HR interview guide:
    Resume:
    {resume_text}
    Company:
    {company_name}
    Role:
    {job_role}
    Generate:
    - 5 key HR / Behavioral Questions (including 'Tell me about yourself', 'Why this company/role')
    - Short, simple, and concise ideal answers for each question.
    """,
        expected_output="HR Interview Guide",
        agent=agent
    )
    crew = Crew(agents=[agent], tasks=[task], verbose=False)
    return str(crew.kickoff())

def run_step6(llm):
    agent = Agent(
        role="Technical Interviewer",
        goal="Generate role-specific technical questions.",
        backstory="Senior Software Engineer.",
        llm=llm,
        verbose=False,
        allow_delegation=False
    )
    task = Task(
        description=f"""
    Generate a short, simple, and concise technical interview guide:
    Resume:
    {resume_text}
    Role:
    {job_role}
    Company:
    {company_name}
    JD:
    {job_description}
    Generate:
    - 5 key technical/coding/project questions (ranging from easy, medium, to hard)
    - Short, simple, and concise expected answers/bullet points for each.
    """,
        expected_output="Technical Interview Guide",
        agent=agent
    )
    crew = Crew(agents=[agent], tasks=[task], verbose=False)
    return str(crew.kickoff())

def run_step7(llm, company_output, resume_output, jd_match_output, resume_enhancement_output, hr_output, technical_output):
    agent = Agent(
        role="Interview Coach",
        goal="Create final interview preparation report.",
        backstory="Career Mentor.",
        llm=llm,
        verbose=False,
        allow_delegation=False
    )
    task = Task(
        description=f"""
    Create a final interview preparation report based on the findings from previous steps.

    Company:
    {company_name}
    Role:
    {job_role}

    Here are the findings from previous analysis steps:
    1. Company Research:
    {company_output}

    2. Resume Analysis & JD Match:
    {resume_output}
    {jd_match_output}

    3. Resume Enhancement Suggestions:
    {resume_enhancement_output}

    4. HR & Technical Preparation:
    {hr_output}
    {technical_output}

    Synthesize these findings and provide:
    1. Overall Readiness Score (0-100)
    2. Top 3 Strengths & 3 Areas to Improve
    3. A simple 30-day learning roadmap
    Keep the final report clean, short, and simple.
    """,
        expected_output="Final Synthesized Interview Report",
        agent=agent
    )
    crew = Crew(agents=[agent], tasks=[task], verbose=False)
    return str(crew.kickoff())

def run_step8(llm):
    agent = Agent(
        role="Mock Interviewer",
        goal="Simulate a real interview.",
        backstory="Senior interviewer.",
        llm=llm,
        verbose=False,
        allow_delegation=False
    )
    task = Task(
        description=f"""
    Simulate a short, simple, and concise mock interview:
    Role:
    {job_role}
    Company:
    {company_name}
    Resume:
    {resume_text}
    Generate:
    - 5 relevant interview questions
    - Short sample answers
    - Simple evaluation criteria
    - Readiness Score (0-100)
    """,
        expected_output="Mock Interview Report",
        agent=agent
    )
    crew = Crew(agents=[agent], tasks=[task], verbose=False)
    return str(crew.kickoff())

def run_step9(llm, resume_enhancement_output):
    agent = Agent(
        role="Resume Creator",
        goal="Create a polished, professional, and copy-paste ready improved sample resume.",
        backstory="Expert Resume Designer & Professional Writer.",
        llm=llm,
        verbose=False,
        allow_delegation=False
    )
    task = Task(
        description=f"""
    Create a complete but short and simple improved sample resume that incorporates the missing keywords and enhancement suggestions:
    Original Resume:
    {resume_text}
    Role:
    {job_role}
    JD:
    {job_description}
    Enhancement Suggestions:
    {resume_enhancement_output}

    Provide a clean, copy-paste ready markdown layout containing:
    - Professional Summary
    - Experience Section (short bullet points)
    - Projects Section (short bullet points)
    - Skills Section
    Keep descriptions short, simple, and concise.
    """,
        expected_output="Improved Sample Resume",
        agent=agent
    )
    crew = Crew(agents=[agent], tasks=[task], verbose=False)
    return str(crew.kickoff())


# Helper to format dynamic status cards
def get_status_card(title, icon, status):
    icon_bg = "linear-gradient(135deg, #e2e8f0, #94a3b8)"  # Fallback
    if "ATS" in title:
        icon_bg = "linear-gradient(135deg, #a7f3d0, #34d399)"
    elif "Research" in title:
        icon_bg = "linear-gradient(135deg, #bfdbfe, #60a5fa)"
    elif "Match" in title:
        icon_bg = "linear-gradient(135deg, #fecdd3, #fb7185)"
    elif "Enhancement" in title or "Expert" in title:
        icon_bg = "linear-gradient(135deg, #fef08a, #facc15)"
    elif "HR" in title:
        icon_bg = "linear-gradient(135deg, #c7d2fe, #818cf8)"
    elif "Technical" in title:
        icon_bg = "linear-gradient(135deg, #a5f3fc, #22d3ee)"
    elif "Mock" in title:
        icon_bg = "linear-gradient(135deg, #e9d5ff, #c084fc)"
    elif "Coach" in title or "Summary" in title:
        icon_bg = "linear-gradient(135deg, #fde047, #ca8a04)"
    elif "Creator" in title or "Resume" in title:
        icon_bg = "linear-gradient(135deg, #e2e8f0, #94a3b8)"

    if status == "pending":
        bg_color = "#121214"
        border_color = "#24242a"
        text_color = "#9ca3af"
        badge = "<span style='color: #9ca3af; background: rgba(255, 255, 255, 0.05); padding: 6px 14px; border-radius: 12px; font-size: 0.8em; font-weight: 600; border: 1px solid rgba(255, 255, 255, 0.12); display: inline-flex; align-items: center; gap: 4px;'>🕒 Pending</span>"
        status_text = "Waiting to start analysis..."
    elif status == "running":
        bg_color = "#0b1523"
        border_color = "#1d3e63"
        text_color = "#ffffff"
        badge = "<span style='color: #60a5fa; background: rgba(59, 130, 246, 0.12); padding: 6px 14px; border-radius: 12px; font-size: 0.8em; font-weight: 600; border: 1px solid rgba(59, 130, 246, 0.3); display: inline-flex; align-items: center; gap: 4px;'>⟳ Working...</span>"
        status_text = "Analyzing data and generating report..."
    elif status == "completed":
        bg_color = "#091913"
        border_color = "#154e38"
        text_color = "#ffffff"
        badge = "<span style='color: #34d399; background: rgba(16, 185, 129, 0.12); padding: 6px 14px; border-radius: 12px; font-size: 0.8em; font-weight: 600; border: 1px solid rgba(16, 185, 129, 0.3); display: inline-flex; align-items: center; gap: 4px;'>✓ Done</span>"
        status_text = "Results compiled successfully!"
    elif status == "failed":
        bg_color = "#211010"
        border_color = "#5c2424"
        text_color = "#ffffff"
        badge = "<span style='color: #f87171; background: rgba(239, 68, 68, 0.12); padding: 6px 14px; border-radius: 12px; font-size: 0.8em; font-weight: 600; border: 1px solid rgba(239, 68, 68, 0.3); display: inline-flex; align-items: center; gap: 4px;'>⚠ Error</span>"
        status_text = "Failed to generate output."

    return f"""
    <div class='agent-card' style='background-color: {bg_color}; border: 1px solid {border_color}; color: {text_color}; font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Helvetica, Arial, sans-serif; transition: all 0.3s ease; box-shadow: 0 10px 25px -5px rgba(0,0,0,0.3); box-sizing: border-box;'>
        <div style='width: 60px; height: 60px; background: {icon_bg}; border-radius: 16px; display: flex; align-items: center; justify-content: center; margin-bottom: 16px; box-shadow: 0 4px 12px rgba(0,0,0,0.15); flex-shrink: 0;'>
            <span style='font-size: 2.0em; line-height: 1;'>{icon}</span>
        </div>
        <div style='font-size: 1.25em; font-weight: 700; color: #ffffff; margin-bottom: 12px; line-height: 1.25; min-height: 48px; display: flex; align-items: center; justify-content: center; width: 100%; text-align: center; word-wrap: break-word; overflow-wrap: break-word;'>
            {title}
        </div>
        <div style='margin-bottom: 16px; display: flex; justify-content: center; width: 100%; flex-shrink: 0;'>
            {badge}
        </div>
        <div style='width: 100%; border-top: 1px solid rgba(255, 255, 255, 0.08); margin: 12px 0; flex-shrink: 0;'></div>
        <div style='font-size: 0.88em; color: #9ca3af; line-height: 1.4; margin-top: auto; width: 100%; text-align: center; word-wrap: break-word; overflow-wrap: break-word;'>
            {status_text}
        </div>
    </div>
    """

card_meta = {
    "resume_output": {"title": "Resume ATS Analyzer", "icon": "📊"},
    "company_output": {"title": "Company Research", "icon": "🔍"},
    "jd_match_output": {"title": "JD Match Analyzer", "icon": "🎯"},
    "resume_enhancement_output": {"title": "Enhancement Expert", "icon": "✍️"},
    "hr_output": {"title": "HR Interviewer", "icon": "🤝"},
    "technical_output": {"title": "Technical Interviewer", "icon": "💻"},
    "mock_interview_output": {"title": "Mock Interview", "icon": "🎙️"},
    "final_output": {"title": "Career Coach Summary", "icon": "🏆"},
    "resume_creation_output": {"title": "Resume Creator", "icon": "📄"},
}

st.write("---")

col_all1, col_all2, col_all3 = st.columns(3)
with col_all1:
    if st.button("🚀 Auto-Generate (Step-by-Step)", type="primary", use_container_width=True):
        # Clear existing cached outputs to start sequential fresh
        for key in [
            "resume_output", "company_output", "jd_match_output", 
            "resume_enhancement_output", "hr_output", "technical_output", 
            "final_output", "mock_interview_output", "resume_creation_output"
        ]:
            if key in st.session_state:
                del st.session_state[key]
        st.session_state.incremental_generating = True
        st.rerun()

with col_all2:
    if st.button("⚡ Generate All in Parallel", type="secondary", use_container_width=True):
        st.session_state.incremental_generating = False
        st.subheader("⚡ Live Multi-Agent Parallel Execution Status")
        # Setup 3x3 grid system matching layout request
        row1_cols = st.columns(3, gap="large")
        row2_cols = st.columns(3, gap="large")
        row3_cols = st.columns(3, gap="large")
        
        status_placeholders = {
            "resume_output": row1_cols[0].empty(),
            "company_output": row1_cols[1].empty(),
            "jd_match_output": row1_cols[2].empty(),
            "resume_enhancement_output": row2_cols[0].empty(),
            "hr_output": row2_cols[1].empty(),
            "technical_output": row2_cols[2].empty(),
            "mock_interview_output": row3_cols[0].empty(),
            "final_output": row3_cols[1].empty(),
            "resume_creation_output": row3_cols[2].empty(),
        }
        
        # Init state UI
        for k, p in status_placeholders.items():
            meta = card_meta[k]
            p.markdown(get_status_card(meta["title"], meta["icon"], "pending"), unsafe_allow_html=True)
            
        try:
            llm = get_llm()
            
            # Phase 1 Parallel Execution: Steps 1, 2, 3, 4, 5, 6
            with ThreadPoolExecutor(max_workers=6) as executor1:
                futures_phase1 = {
                    executor1.submit(run_step1, llm): "resume_output",
                    executor1.submit(run_step2, llm): "company_output",
                    executor1.submit(run_step3, llm): "jd_match_output",
                    executor1.submit(run_step4, llm): "resume_enhancement_output",
                    executor1.submit(run_step5, llm): "hr_output",
                    executor1.submit(run_step6, llm): "technical_output",
                }
                
                results = {}
                has_error = False
                error_msg = None
                
                while not all(f.done() for f in futures_phase1.keys()):
                    for future, key in futures_phase1.items():
                        meta = card_meta[key]
                        if future.done():
                            try:
                                if key not in results:
                                    results[key] = future.result()
                                status = "completed"
                            except Exception as e:
                                status = "failed"
                                has_error = True
                                error_msg = e
                        elif future.running():
                            status = "running"
                        else:
                            status = "pending"
                        status_placeholders[key].markdown(get_status_card(meta["title"], meta["icon"], status), unsafe_allow_html=True)
                    
                    # Phase 2 tasks (Mock Interview, Career Coach, Resume Creator) remain pending in Phase 1
                    for key in ["final_output", "mock_interview_output", "resume_creation_output"]:
                        meta = card_meta[key]
                        status_placeholders[key].markdown(get_status_card(meta["title"], meta["icon"], "pending"), unsafe_allow_html=True)
                        
                    if has_error:
                        break
                    time.sleep(0.25)
                    
                if not has_error:
                    for future, key in futures_phase1.items():
                        meta = card_meta[key]
                        if key not in results:
                            results[key] = future.result()
                        status_placeholders[key].markdown(get_status_card(meta["title"], meta["icon"], "completed"), unsafe_allow_html=True)
                        
            if has_error:
                raise error_msg
                
            # Phase 2 Parallel Execution: Steps 7, 8, 9 (Career Coach, Mock Interview, Resume Creator)
            with ThreadPoolExecutor(max_workers=3) as executor2:
                futures_phase2 = {
                    executor2.submit(run_step7, llm, results["company_output"], results["resume_output"], results["jd_match_output"], results["resume_enhancement_output"], results["hr_output"], results["technical_output"]): "final_output",
                    executor2.submit(run_step8, llm): "mock_interview_output",
                    executor2.submit(run_step9, llm, results["resume_enhancement_output"]): "resume_creation_output",
                }
                
                while not all(f.done() for f in futures_phase2.keys()):
                    for future, key in futures_phase2.items():
                        meta = card_meta[key]
                        if future.done():
                            try:
                                if key not in results:
                                    results[key] = future.result()
                                status = "completed"
                            except Exception as e:
                                status = "failed"
                                has_error = True
                                error_msg = e
                        elif future.running():
                            status = "running"
                        else:
                            status = "pending"
                        status_placeholders[key].markdown(get_status_card(meta["title"], meta["icon"], status), unsafe_allow_html=True)
                    
                    if has_error:
                        break
                    time.sleep(0.25)
                    
                if not has_error:
                    for future, key in futures_phase2.items():
                        meta = card_meta[key]
                        if key not in results:
                            results[key] = future.result()
                        status_placeholders[key].markdown(get_status_card(meta["title"], meta["icon"], "completed"), unsafe_allow_html=True)
                        
            if has_error:
                raise error_msg
                
            # Store in session state
            for key, val in results.items():
                st.session_state[key] = val
                
            st.success("✅ Generated all reports successfully!")
            st.rerun()
            
        except Exception as exc:
            if is_quota_exhausted_error(exc):
                st.error("❌ Your OpenRouter API key rate limit or credits are exhausted. Please check your account.")
            else:
                st.error("Failed to generate all reports.")
                st.exception(exc)
            st.stop()

with col_all3:
    if st.button("🗑️ Reset & Clear All Cached Reports", type="secondary", use_container_width=True):
        st.session_state.incremental_generating = False
        for key in [
            "resume_output", "company_output", "jd_match_output", 
            "resume_enhancement_output", "hr_output", "technical_output", 
            "final_output", "mock_interview_output", "resume_creation_output"
        ]:
            if key in st.session_state:
                del st.session_state[key]
        st.rerun()

st.subheader("🛠️ Step-by-Step Interview Preparation Modules")

# Define all steps configs
steps_config = [
    {
        "key": "resume_output",
        "title": "📊 Step 1: ATS Resume Analysis",
        "file_name": "ATS_Resume_Analysis.txt",
        "runner": lambda llm: run_step1(llm),
        "deps": []
    },
    {
        "key": "company_output",
        "title": "🔍 Step 2: Company Research",
        "file_name": "Company_Research.txt",
        "runner": lambda llm: run_step2(llm),
        "deps": []
    },
    {
        "key": "jd_match_output",
        "title": "🎯 Step 3: Job Description Match Analysis",
        "file_name": "JD_Match_Analysis.txt",
        "runner": lambda llm: run_step3(llm),
        "deps": []
    },
    {
        "key": "resume_enhancement_output",
        "title": "✍️ Step 4: Resume Enhancement Suggestions",
        "file_name": "Resume_Enhancement_Suggestions.txt",
        "runner": lambda llm: run_step4(llm),
        "deps": []
    },
    {
        "key": "hr_output",
        "title": "🤝 Step 5: HR Interview Guide",
        "file_name": "HR_Interview_Guide.txt",
        "runner": lambda llm: run_step5(llm),
        "deps": []
    },
    {
        "key": "technical_output",
        "title": "💻 Step 6: Technical Interview Guide",
        "file_name": "Technical_Interview_Guide.txt",
        "runner": lambda llm: run_step6(llm),
        "deps": []
    },
    {
        "key": "final_output",
        "title": "🏆 Step 7: Final Career Coach Summary & 30-Day Plan",
        "file_name": "Final_Career_Coach_Summary.txt",
        "runner": lambda llm: run_step7(
            llm,
            st.session_state.company_output, 
            st.session_state.resume_output, 
            st.session_state.jd_match_output, 
            st.session_state.resume_enhancement_output, 
            st.session_state.hr_output, 
            st.session_state.technical_output
        ),
        "deps": ["company_output", "resume_output", "jd_match_output", "resume_enhancement_output", "hr_output", "technical_output"]
    },
    {
        "key": "mock_interview_output",
        "title": "🎙️ Step 8: Mock Interview Simulator",
        "file_name": "Mock_Interview_Simulator.txt",
        "runner": lambda llm: run_step8(llm),
        "deps": []
    },
    {
        "key": "resume_creation_output",
        "title": "📄 Step 9: Improved Sample Resume",
        "file_name": "Improved_Sample_Resume.txt",
        "runner": lambda llm: run_step9(llm, st.session_state.resume_enhancement_output),
        "deps": ["resume_enhancement_output"]
    }
]

# Find current running step in auto-generation mode for localized status indicators
current_generating_key = None
if st.session_state.get("incremental_generating", False):
    for step in steps_config:
        if step["key"] not in st.session_state:
            current_generating_key = step["key"]
            break

# Render each step
for step in steps_config:
    key = step["key"]
    title = step["title"]
    file_name = step["file_name"]
    runner = step["runner"]
    deps = step["deps"]
    
    with st.expander(title, expanded=(key in st.session_state or key == current_generating_key)):
        if key in st.session_state:
            # Display report content
            st.markdown(st.session_state[key])
            st.write("---")
            
            # Action buttons for this report
            c1, c2 = st.columns(2)
            with c1:
                st.download_button(
                    label=f"Download {card_meta[key]['title']} Report",
                    data=st.session_state[key],
                    file_name=file_name,
                    mime="text/plain",
                    use_container_width=True,
                    key=f"dl_{key}"
                )
            with c2:
                if st.button(f"🔄 Regenerate This Module", key=f"regen_{key}", use_container_width=True):
                    try:
                        with st.spinner(f"Regenerating {card_meta[key]['title']}..."):
                            llm = get_llm()
                            st.session_state[key] = runner(llm)
                        st.success(f"Regenerated {card_meta[key]['title']} successfully!")
                        st.rerun()
                    except Exception as e:
                        if is_quota_exhausted_error(e):
                            st.error("❌ Your OpenRouter API key rate limit or credits are exhausted.")
                        else:
                            st.error("Failed to regenerate.")
                            st.exception(e)
        else:
            if key == current_generating_key:
                st.info("🔄 **Auto-generating this report now... Please wait.**")
                st.spinner("Executing agent...")
            else:
                # Report not generated yet
                st.info("Report not generated yet.")
                
                # Check dependencies
                missing_deps = [card_meta[d]["title"] for d in deps if d not in st.session_state]
                if missing_deps:
                    st.warning(f"Prerequisite modules required before running this: {', '.join(missing_deps)}")
                    st.button(f"Generate {card_meta[key]['title']}", key=f"gen_{key}", disabled=True, use_container_width=True)
                else:
                    if st.button(f"Generate {card_meta[key]['title']}", key=f"gen_{key}", type="secondary", use_container_width=True):
                        try:
                            with st.spinner(f"Generating {card_meta[key]['title']}..."):
                                llm = get_llm()
                                st.session_state[key] = runner(llm)
                            st.success(f"Generated {card_meta[key]['title']} successfully!")
                            st.rerun()
                        except Exception as e:
                            if is_quota_exhausted_error(e):
                                st.error("❌ Your OpenRouter API key rate limit or credits are exhausted.")
                            else:
                                st.error("Failed to generate report.")
                                st.exception(e)

# Combined Report Assembly (Only includes reports that have been generated so far)
st.write("---")
st.subheader("🎉 Combined Interview Preparation Report Preview")

available_reports = []
for step in steps_config:
    k = step["key"]
    title_label = card_meta[k]["title"].upper()
    if k in st.session_state:
        available_reports.append(f"""=========================================
{title_label}
=========================================
{st.session_state[k]}
""")

if available_reports:
    combined_report = f"""# INTERVIEW PREPARATION REPORT
Company: {company_name}
Role: {job_role}

""" + "\n\n".join(available_reports)

    st.text_area("Full Compiled Report Preview (generated modules only)", combined_report, height=350)
    
    report_path = get_temp_storage_dir() / "Interview_Preparation_Report.txt"
    try:
        with open(report_path, "w", encoding="utf-8") as f:
            f.write(combined_report)
    except Exception:
        report_path = None
        
    col1, col2 = st.columns(2)
    with col1:
        st.download_button(
            label="Download Combined Report",
            data=combined_report,
            file_name="Interview_Preparation_Report.txt",
            mime="text/plain",
            use_container_width=True
        )
    with col2:
        if st.button("🗑️ Reset All Reports", type="secondary", use_container_width=True, key="reset_bottom"):
            for key in [
                "resume_output", "company_output", "jd_match_output", 
                "resume_enhancement_output", "hr_output", "technical_output", 
                "final_output", "mock_interview_output", "resume_creation_output"
            ]:
                if key in st.session_state:
                    del st.session_state[key]
            st.rerun()
    if report_path:
        st.success(f"Combined report saved to {report_path}")
    else:
        st.success("Combined report generated successfully. Use the download button to save it.")
else:
    st.info("No reports generated yet. Click 'Auto-Generate (Step-by-Step)' above or generate individual modules separately.")

# Sequential Incremental Auto-Generator Executor Block
if st.session_state.get("incremental_generating", False):
    next_step = None
    for step in steps_config:
        if step["key"] not in st.session_state:
            next_step = step
            break
            
    if next_step:
        key = next_step["key"]
        runner = next_step["runner"]
        deps = next_step["deps"]
        
        # Check prerequisites
        missing_deps = [d for d in deps if d not in st.session_state]
        if missing_deps:
            st.session_state.incremental_generating = False
            st.error(f"Cannot proceed: Prerequisite modules are missing for {card_meta[key]['title']}.")
            st.stop()
            
        step_idx = steps_config.index(next_step) + 1
        
        # Render a localized loader at the bottom of the page
        st.write("---")
        with st.spinner(f"⏳ **Auto-Generating Module {step_idx}/9: {card_meta[key]['title']}**... Please wait."):
            try:
                llm = get_llm()
                st.session_state[key] = runner(llm)
                st.rerun()
            except Exception as e:
                st.session_state.incremental_generating = False
                if is_quota_exhausted_error(e):
                    st.error(f"❌ OpenRouter API rate limit or quota exceeded during generation of {card_meta[key]['title']}.")
                else:
                    st.error(f"An error occurred while generating {card_meta[key]['title']}.")
                    st.exception(e)
                st.stop()
    else:
        st.session_state.incremental_generating = False
        st.success("🎉 All interview preparation reports have been generated successfully!")
        st.rerun()